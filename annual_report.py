from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from matcher import normalize_text, parse_amount
from month_state import months_dir

MONTHS = {
    "styczen": 1,
    "luty": 2,
    "marzec": 3,
    "kwiecien": 4,
    "maj": 5,
    "czerwiec": 6,
    "lipiec": 7,
    "sierpien": 8,
    "wrzesien": 9,
    "pazdziernik": 10,
    "listopad": 11,
    "grudzien": 12,
}

MONTH_NAMES = {
    1: "Styczeń", 2: "Luty", 3: "Marzec", 4: "Kwiecień",
    5: "Maj", 6: "Czerwiec", 7: "Lipiec", 8: "Sierpień",
    9: "Wrzesień", 10: "Październik", 11: "Listopad", 12: "Grudzień",
}


def _month_year(label: str) -> tuple[int | None, int | None]:
    text = normalize_text(label)
    month = next((number for name, number in MONTHS.items() if name in text), None)
    match = re.search(r"(?<!\d)(20\d{2}|\d{2})(?!\d)", text)
    year = None
    if match:
        raw = int(match.group(1))
        year = raw if raw >= 2000 else 2000 + raw
    return month, year


def _states(root: str | Path | None = None) -> list[dict]:
    base = Path(root) if root else months_dir()
    if not base.exists():
        return []
    output = []
    for path in base.glob("*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(data, dict):
            output.append(data)
    return output


def available_years(root: str | Path | None = None) -> list[int]:
    years = set()
    for state in _states(root):
        _month, year = _month_year(str(state.get("month_label") or state.get("budget_sheet") or ""))
        if year:
            years.add(year)
    return sorted(years, reverse=True)


def annual_report(year: int, root: str | Path | None = None) -> dict:
    monthly: dict[int, dict] = {}
    categories: defaultdict[str, float] = defaultdict(float)
    annual_income = 0.0
    annual_expenses = 0.0
    annual_paid = 0.0
    annual_review = 0.0
    annual_missing = 0.0

    for state in _states(root):
        label = str(state.get("month_label") or state.get("budget_sheet") or "")
        month, state_year = _month_year(label)
        if not month or state_year != year:
            continue

        items = list(state.get("items") or [])
        results = list(state.get("results") or [])
        income = 0.0
        expenses = 0.0
        for item in items:
            amount = abs(parse_amount(item.get("amount")) or 0.0)
            if item.get("entry_type") == "income":
                income += amount
            else:
                expenses += amount
                category = str(item.get("category") or "Inne").strip() or "Inne"
                categories[category] += amount

        paid = review = missing = 0.0
        paid_count = review_count = missing_count = 0
        for row in results:
            if row.get("entry_type") == "income":
                continue
            amount = abs(parse_amount(row.get("amount")) or 0.0)
            status = str(row.get("status") or "")
            if status == "OPŁACONA":
                paid += amount
                paid_count += 1
            elif status == "DO SPRAWDZENIA":
                review += amount
                review_count += 1
            elif status == "BRAK":
                missing += amount
                missing_count += 1

        monthly[month] = {
            "month": month,
            "month_name": MONTH_NAMES[month],
            "label": label,
            "income": round(income, 2),
            "expenses": round(expenses, 2),
            "paid": round(paid, 2),
            "review": round(review, 2),
            "missing": round(missing, 2),
            "balance": round(income - expenses, 2),
            "paid_count": paid_count,
            "review_count": review_count,
            "missing_count": missing_count,
        }
        annual_income += income
        annual_expenses += expenses
        annual_paid += paid
        annual_review += review
        annual_missing += missing

    months = [monthly[key] for key in sorted(monthly)]
    count = len(months)
    avg_expenses = annual_expenses / count if count else 0.0
    avg_income = annual_income / count if count else 0.0
    best = max(months, key=lambda row: row["balance"], default=None)
    most_expensive = max(months, key=lambda row: row["expenses"], default=None)
    cheapest = min(months, key=lambda row: row["expenses"], default=None)

    category_rows = [
        {"category": name, "amount": round(amount, 2), "share": round((amount / annual_expenses * 100) if annual_expenses else 0.0, 1)}
        for name, amount in categories.items()
    ]
    category_rows.sort(key=lambda row: row["amount"], reverse=True)

    insights = []
    if most_expensive:
        insights.append(f"Najdroższy miesiąc: {most_expensive['month_name']} — {most_expensive['expenses']:.2f} zł planowanych wydatków.")
    if cheapest and count > 1:
        insights.append(f"Najtańszy miesiąc: {cheapest['month_name']} — {cheapest['expenses']:.2f} zł.")
    if category_rows:
        top = category_rows[0]
        insights.append(f"Największa kategoria: {top['category']} — {top['amount']:.2f} zł ({top['share']:.1f}% wydatków).")
    if best:
        insights.append(f"Najlepszy planowany bilans: {best['month_name']} — {best['balance']:+.2f} zł.")
    if annual_review or annual_missing:
        insights.append(f"Do uporządkowania w zapisanych miesiącach: {annual_review + annual_missing:.2f} zł.")

    return {
        "year": int(year),
        "months_count": count,
        "income": round(annual_income, 2),
        "expenses": round(annual_expenses, 2),
        "paid": round(annual_paid, 2),
        "review": round(annual_review, 2),
        "missing": round(annual_missing, 2),
        "balance": round(annual_income - annual_expenses, 2),
        "average_income": round(avg_income, 2),
        "average_expenses": round(avg_expenses, 2),
        "months": months,
        "categories": category_rows,
        "insights": insights,
    }


def export_annual_report_xlsx(report: dict, target: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = Path(target)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Rok {report['year']}"

    ws.append(["PayCheck — podsumowanie roczne", report["year"]])
    ws["A1"].font = Font(bold=True, size=16)
    ws.append([])
    summary = (
        ("Wpływy", report["income"]),
        ("Plan wydatków", report["expenses"]),
        ("Potwierdzone", report["paid"]),
        ("Do sprawdzenia", report["review"]),
        ("Brak potwierdzenia", report["missing"]),
        ("Bilans planu", report["balance"]),
        ("Średnie wpływy / mies.", report["average_income"]),
        ("Średnie wydatki / mies.", report["average_expenses"]),
    )
    for label, value in summary:
        ws.append([label, value])
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    month_ws = wb.create_sheet("Miesiące")
    month_ws.append(["Miesiąc", "Wpływy", "Wydatki", "Potwierdzone", "Do sprawdzenia", "Brak", "Bilans"])
    for row in report["months"]:
        month_ws.append([row["month_name"], row["income"], row["expenses"], row["paid"], row["review"], row["missing"], row["balance"]])

    cat_ws = wb.create_sheet("Kategorie")
    cat_ws.append(["Kategoria", "Kwota", "Udział %"])
    for row in report["categories"]:
        cat_ws.append([row["category"], row["amount"], row["share"]])

    insight_ws = wb.create_sheet("Wnioski")
    insight_ws.append(["Najważniejsze wnioski"])
    insight_ws["A1"].font = Font(bold=True, size=14)
    for text in report["insights"]:
        insight_ws.append([text])

    for sheet in wb.worksheets:
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
            sheet.column_dimensions[column[0].column_letter].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
