from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


SUMMARY_PREFIXES = (
    "suma ",
    "koszt mies",
    "800+",
    "kasa dodatkowa",
    "wyplata",
    "wypłata",
    "reszta",
    "saldo po",
    "pieniadze od",
    "pieniądze od",
    "dlug u",
    "dług u",
)


def list_budget_sheets(path: str | Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _is_budget_sheet(ws) -> bool:
    headers = [str(ws.cell(1, col).value or "").strip().lower() for col in range(1, 5)]
    return (
        len(headers) >= 4
        and headers[1] == "instytucja"
        and headers[2] == "suma"
        and headers[3] == "miesiecznie"
    )


def load_budget_sheet(path: str | Path, sheet_name: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Nie ma arkusza: {sheet_name}")
        ws = wb[sheet_name]
        if not _is_budget_sheet(ws):
            raise ValueError(
                "Wybrany arkusz nie ma układu budżetu: instytucja / Suma / Miesiecznie."
            )

        items: list[dict] = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
            end_date, institution, name, monthly = row
            name_text = str(name or "").strip()
            if not name_text or not isinstance(monthly, (int, float)) or monthly <= 0:
                continue

            normalized = " ".join(name_text.lower().split())
            if any(normalized.startswith(prefix) for prefix in SUMMARY_PREFIXES):
                continue

            institution_text = str(institution or "").strip()
            search_name = " ".join(part for part in (institution_text, name_text) if part)
            items.append(
                {
                    "invoice_no": f"BUD-{row_no}",
                    "counterparty": search_name or name_text,
                    "display_name": name_text,
                    "institution": institution_text,
                    "amount": float(monthly),
                    # W arkuszu budżetu kolumna A oznacza datę zakończenia raty/umowy,
                    # a nie termin płatności w wybranym miesiącu. Nie używamy jej do matchingu.
                    "date": "",
                    "end_date": end_date if end_date else "",
                    "source_type": "budget",
                    "source_sheet": sheet_name,
                    "source_row": row_no,
                }
            )

        if not items:
            raise ValueError("W wybranym arkuszu nie znaleziono miesięcznych kosztów do porównania.")
        return items
    finally:
        wb.close()
