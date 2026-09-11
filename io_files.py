from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from bank_formats_320 import parse_mt940


INVOICE_ALIASES = {
    "counterparty": ("kontrahent", "odbiorca", "firma", "nazwa"),
    "amount": ("kwota", "do zaplaty", "do zapłaty", "wartosc", "wartość"),
    "date": ("termin", "data platnosci", "data płatności", "data"),
    "invoice_no": ("nr faktury", "numer faktury", "faktura", "numer"),
}

BANK_ALIASES = {
    "counterparty": (
        "odbiorca", "kontrahent", "nadawca", "nazwa", "opis kontrahenta",
        "nazwa kontrahenta", "beneficjent", "zleceniodawca",
    ),
    "amount": (
        "kwota", "wartosc", "wartość", "kwota operacji", "kwota transakcji",
        "kwota w walucie rachunku", "obroty",
    ),
    "date": (
        "data operacji", "data ksiegowania", "data księgowania", "data",
        "data transakcji", "data waluty",
    ),
    "title": (
        "tytul", "tytuł", "opis", "tytul operacji", "tytuł operacji",
        "opis operacji", "szczegóły", "szczegoly",
    ),
}

_DATE_RE = re.compile(r"^\s*(\d{2}\.\d{2}\.\d{4})\b")
_AMOUNT_RE = re.compile(r"(?<!\d)([-+]?\s*\d{1,3}(?:[ .]\d{3})*,\d{2})\s*PLN\b", re.IGNORECASE)


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _map_headers(headers: list[object], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    normalized = [_normalize_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for key, choices in aliases.items():
        for idx, header in enumerate(normalized):
            if header in choices:
                mapping[key] = idx
                break
    return mapping


def _rows_from_xlsx(path: str | Path) -> tuple[list[object], list[list[object]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration:
        return [], []
    return headers, [list(row) for row in iterator]


def _rows_from_csv(path: str | Path) -> tuple[list[str], list[list[str]]]:
    raw = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Nie udało się odczytać kodowania pliku CSV.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        rows = list(csv.reader(text.splitlines(), dialect))
    except csv.Error:
        rows = list(csv.reader(text.splitlines(), delimiter=";"))

    if not rows:
        return [], []
    return rows[0], rows[1:]


def _parse_pln_amount(value: str) -> float:
    normalized = value.replace(" ", "").replace(".", "").replace(",", ".")
    return float(normalized)


def _extract_counterparty(description: str) -> str:
    for label in ("Odbiorca:", "Nadawca:"):
        match = re.search(
            rf"{label}\s*(.+?)(?=\s+(?:Tytuł:|Tytul:|Nr rachunku:|Rachunek:)|$)",
            description,
            flags=re.IGNORECASE,
        )
        if match:
            return " ".join(match.group(1).split()).strip(" ,;-")

    card = re.search(
        r"\bw\s+(.+?)(?=,\s*[A-ZĄĆĘŁŃÓŚŹŻ][A-ZĄĆĘŁŃÓŚŹŻ .-]{1,30}(?:,\s*[A-Z]{2})?\b|$)",
        description,
        flags=re.IGNORECASE,
    )
    if card:
        return " ".join(card.group(1).split()).strip(" ,;-")

    return description[:120].strip()


def _extract_title(description: str) -> str:
    match = re.search(r"(?:Tytuł|Tytul):\s*(.+)$", description, flags=re.IGNORECASE)
    if match:
        return " ".join(match.group(1).split())
    return description


def _parse_bank_pdf_text(text: str) -> list[dict]:
    lines = [" ".join(line.split()) for line in text.splitlines() if line.strip()]
    blocks: list[str] = []
    current: list[str] = []

    for line in lines:
        if _DATE_RE.match(line):
            if current:
                blocks.append(" ".join(current))
            current = [line]
        elif current:
            current.append(line)

    if current:
        blocks.append(" ".join(current))

    transactions: list[dict] = []
    for block in blocks:
        date_match = _DATE_RE.match(block)
        if not date_match:
            continue

        transaction_date = date_match.group(1)
        rest = block[date_match.end():].strip()

        booking_match = re.match(r"^(?:-|\d{2}\.\d{2}\.\d{4})\b", rest)
        if booking_match:
            rest = rest[booking_match.end():].strip()

        amount_matches = list(_AMOUNT_RE.finditer(rest))
        if not amount_matches:
            continue

        # W opisach kartowych bank potrafi umieścić kwotę informacyjną, a dopiero
        # na końcu właściwą zaksięgowaną kwotę ze znakiem. Bierzemy ostatnią.
        amount_match = amount_matches[-1]
        description = rest[:amount_match.start()].strip(" -|;")
        if not description:
            continue

        amount = _parse_pln_amount(amount_match.group(1))
        transactions.append(
            {
                "date": transaction_date,
                "counterparty": _extract_counterparty(description),
                "amount": amount,
                "title": _extract_title(description),
                "description": description,
                "source": "PDF",
            }
        )

    if not transactions:
        raise ValueError(
            "Nie znaleziono transakcji w PDF. Jeśli to skan/zdjęcie bez warstwy tekstowej, "
            "wyeksportuj wyciąg jako PDF tekstowy, CSV, XLSX albo MT940/STA."
        )
    return transactions


def _rows_from_pdf(path: str | Path) -> list[dict]:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        raise ValueError(
            "PDF nie zawiera tekstu do odczytu. Prawdopodobnie jest skanem; "
            "użyj PDF tekstowego, CSV, XLSX albo MT940/STA."
        )
    return _parse_bank_pdf_text(text)


def _read_table(path: str | Path) -> tuple[list[object], list[list[object]]]:
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return _rows_from_xlsx(path)
    if suffix == ".csv":
        return _rows_from_csv(path)
    raise ValueError("Obsługiwane formaty tabel: XLSX i CSV.")


def _convert_rows(
    headers: list[object],
    rows: Iterable[list[object]],
    aliases: dict[str, tuple[str, ...]],
    required: tuple[str, ...],
) -> list[dict]:
    mapping = _map_headers(headers, aliases)
    missing = [key for key in required if key not in mapping]
    if missing:
        raise ValueError("Brak wymaganych kolumn: " + ", ".join(missing))

    converted: list[dict] = []
    for row in rows:
        if not any(value not in (None, "") for value in row):
            continue
        item = {}
        for key, idx in mapping.items():
            item[key] = row[idx] if idx < len(row) else ""
        converted.append(item)
    return converted


def load_invoices(path: str | Path) -> list[dict]:
    headers, rows = _read_table(path)
    return _convert_rows(
        headers,
        rows,
        INVOICE_ALIASES,
        ("counterparty", "amount", "date"),
    )


def load_bank_statement(path: str | Path) -> list[dict]:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return _rows_from_pdf(path)
    if suffix in {".sta", ".mt940", ".940", ".txt"}:
        return parse_mt940(path)

    headers, rows = _read_table(path)
    return _convert_rows(
        headers,
        rows,
        BANK_ALIASES,
        ("counterparty", "amount", "date"),
    )


def save_results(path: str | Path, results: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wynik"

    columns = [
        ("invoice_no", "Nr faktury"),
        ("counterparty", "Kontrahent"),
        ("date", "Termin"),
        ("amount", "Kwota z Excela"),
        ("status", "Status"),
        ("match_score", "Zgodność %"),
        ("bank_date", "Data z banku"),
        ("bank_amount", "Kwota z banku"),
        ("amount_diff", "Różnica kwoty"),
        ("days_diff", "Różnica dni"),
        ("bank_counterparty", "Kontrahent z banku"),
        ("bank_title", "Tytuł przelewu"),
    ]
    ws.append([label for _, label in columns])
    for result in results:
        ws.append([result.get(key, "") for key, _ in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        width = min(
            45,
            max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2),
        )
        ws.column_dimensions[column_cells[0].column_letter].width = width

    wb.save(path)
