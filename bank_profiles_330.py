from __future__ import annotations

import csv
import json
import os
import re
import uuid
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from matcher import normalize_text, parse_amount, parse_date


BUILTIN_PROFILES = [
    {
        "id": "builtin-generic-pl",
        "name": "Polski bank — profil ogólny",
        "bank": "Ogólny",
        "headers": {
            "date": ["data operacji", "data księgowania", "data ksiegowania", "data"],
            "amount": ["kwota", "kwota operacji", "wartość", "wartosc"],
            "counterparty": ["kontrahent", "odbiorca", "nadawca", "opis kontrahenta", "nazwa"],
            "title": ["tytuł", "tytul", "opis", "tytuł operacji", "tytul operacji"],
            "type": ["typ operacji", "rodzaj operacji", "typ transakcji"],
            "balance": ["saldo", "saldo po operacji"],
            "currency": ["waluta"],
        },
    },
    {
        "id": "builtin-velobank",
        "name": "VeloBank / Getin — profil startowy",
        "bank": "VeloBank",
        "headers": {
            "date": ["data operacji", "data księgowania", "data ksiegowania", "data"],
            "amount": ["kwota", "kwota operacji"],
            "counterparty": ["nadawca/odbiorca", "nadawca odbiorca", "kontrahent", "odbiorca", "nadawca"],
            "title": ["tytuł operacji", "tytul operacji", "opis operacji", "opis", "tytuł", "tytul"],
            "type": ["typ operacji", "rodzaj operacji"],
            "balance": ["saldo po operacji", "saldo"],
            "currency": ["waluta"],
        },
    },
    {
        "id": "builtin-pko",
        "name": "PKO BP — profil startowy",
        "bank": "PKO BP",
        "headers": {
            "date": ["data operacji", "data waluty", "data"],
            "amount": ["kwota", "kwota operacji"],
            "counterparty": ["nadawca/odbiorca", "kontrahent", "odbiorca", "nadawca"],
            "title": ["tytuł", "tytul", "opis operacji", "szczegóły", "szczegoly"],
            "type": ["typ operacji", "rodzaj operacji"],
            "balance": ["saldo po operacji", "saldo"],
            "currency": ["waluta"],
        },
    },
    {
        "id": "builtin-mbank",
        "name": "mBank — profil startowy",
        "bank": "mBank",
        "headers": {
            "date": ["data operacji", "data księgowania", "data"],
            "amount": ["kwota", "kwota operacji"],
            "counterparty": ["odbiorca", "nadawca", "kontrahent", "nazwa kontrahenta"],
            "title": ["opis operacji", "opis", "tytuł", "tytul"],
            "type": ["kategoria", "typ operacji", "rodzaj operacji"],
            "balance": ["saldo po operacji", "saldo"],
            "currency": ["waluta"],
        },
    },
]

REQUIRED_FIELDS = ("date", "amount")
OPTIONAL_FIELDS = ("counterparty", "title", "type", "balance", "currency")


def app_dir() -> Path:
    return Path(os.getenv("APPDATA") or Path.home()) / "PayCheck"


def custom_profiles_path() -> Path:
    return app_dir() / "bank_profiles.json"


def _norm(value) -> str:
    return normalize_text(str(value or ""))


def load_custom_profiles() -> list[dict]:
    path = custom_profiles_path()
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    return data if isinstance(data, list) else []


def save_custom_profiles(rows: list[dict]) -> None:
    path = custom_profiles_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def all_profiles() -> list[dict]:
    return load_custom_profiles() + BUILTIN_PROFILES


def save_learned_profile(name: str, mapping: dict[str, str], bank: str = "Własny profil") -> dict:
    clean = {k: str(v).strip() for k, v in mapping.items() if str(v).strip()}
    if not all(field in clean for field in REQUIRED_FIELDS):
        raise ValueError("Profil musi wskazywać kolumnę daty i kwoty.")
    rows = load_custom_profiles()
    profile = {
        "id": f"profile-{uuid.uuid4().hex[:10]}",
        "name": str(name or "Mój profil banku").strip(),
        "bank": str(bank or "Własny profil").strip(),
        "mapping": clean,
    }
    rows.append(profile)
    save_custom_profiles(rows)
    return profile


def delete_learned_profile(profile_id: str) -> bool:
    rows = load_custom_profiles()
    new_rows = [row for row in rows if row.get("id") != profile_id]
    if len(new_rows) == len(rows):
        return False
    save_custom_profiles(new_rows)
    return True


def _read_csv_preview(path: str | Path) -> tuple[list[str], list[list[str]]]:
    raw = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Nie udało się odczytać pliku CSV.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        rows = list(csv.reader(text.splitlines(), dialect))
    except csv.Error:
        rows = list(csv.reader(text.splitlines(), delimiter=";"))
    return (rows[0], rows[1:]) if rows else ([], [])


def _read_xlsx_preview(path: str | Path) -> tuple[list[object], list[list[object]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration:
        return [], []
    rows = []
    for index, row in enumerate(iterator):
        rows.append(list(row))
        if index >= 199:
            break
    return headers, rows


def read_table_preview(path: str | Path) -> tuple[list[str], list[list[object]]]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        headers, rows = _read_csv_preview(path)
    elif suffix == ".xlsx":
        headers, rows = _read_xlsx_preview(path)
    else:
        raise ValueError("Kreator profilu obsługuje CSV i XLSX.")
    return [str(h or "").strip() for h in headers], rows


def _profile_mapping(profile: dict, headers: list[str]) -> dict[str, str]:
    if profile.get("mapping"):
        return {k: v for k, v in profile["mapping"].items() if v in headers}
    normalized = {_norm(header): header for header in headers}
    mapping = {}
    for field, aliases in profile.get("headers", {}).items():
        for alias in aliases:
            if _norm(alias) in normalized:
                mapping[field] = normalized[_norm(alias)]
                break
    return mapping


def detect_profile(headers: list[str]) -> dict:
    best_profile = None
    best_mapping = {}
    best_score = -1
    for profile in all_profiles():
        mapping = _profile_mapping(profile, headers)
        required_hits = sum(field in mapping for field in REQUIRED_FIELDS)
        optional_hits = sum(field in mapping for field in OPTIONAL_FIELDS)
        score = required_hits * 35 + optional_hits * 6
        if required_hits == len(REQUIRED_FIELDS):
            score += 10
        if score > best_score:
            best_profile = profile
            best_mapping = mapping
            best_score = score
    confidence = max(0, min(100, best_score)) if best_profile else 0
    return {
        "profile": best_profile,
        "mapping": best_mapping,
        "confidence": confidence,
        "ready": all(field in best_mapping for field in REQUIRED_FIELDS),
    }


def classify_payment_type(*values) -> str:
    text = _norm(" ".join(str(v or "") for v in values))
    rules = (
        ("BLIK", ("blik",)),
        ("KARTA", ("karta", "card", "visa", "mastercard")),
        ("POLECENIE ZAPŁATY", ("polecenie zaplaty", "direct debit")),
        ("RATA / KREDYT", ("rata", "kredyt", "pozyczka", "pożyczka")),
        ("WYPŁATA GOTÓWKI", ("bankomat", "wyplata gotowki", "wypłata gotówki", "atm")),
        ("WPŁATA GOTÓWKI", ("wplata gotowki", "wpłata gotówki")),
        ("PRZELEW", ("przelew", "transfer")),
    )
    for label, phrases in rules:
        if any(_norm(phrase) in text for phrase in phrases):
            return label
    return "INNE"


def _looks_internal(*values) -> bool:
    text = _norm(" ".join(str(v or "") for v in values))
    phrases = (
        "przelew wlasny", "przelew własny", "miedzy rachunkami", "między rachunkami",
        "transfer wlasny", "transfer własny", "rachunek wlasny", "rachunek własny",
    )
    return any(_norm(phrase) in text for phrase in phrases)


def convert_with_mapping(
    headers: list[str],
    rows: Iterable[list[object]],
    mapping: dict[str, str],
    *,
    profile_name: str,
    source_file: str = "",
) -> list[dict]:
    indexes = {header: index for index, header in enumerate(headers)}
    converted = []
    for raw_row in rows:
        if not any(value not in (None, "") for value in raw_row):
            continue
        raw = {headers[i]: raw_row[i] if i < len(raw_row) else "" for i in range(len(headers))}
        item = {}
        for field, header in mapping.items():
            if header in indexes:
                item[field] = raw_row[indexes[header]] if indexes[header] < len(raw_row) else ""
        if not item.get("date") or item.get("amount") in (None, ""):
            continue
        description = item.get("title") or item.get("counterparty") or ""
        item["payment_type"] = classify_payment_type(item.get("type"), item.get("title"), item.get("counterparty"))
        item["is_internal_transfer"] = _looks_internal(item.get("type"), item.get("title"), item.get("counterparty"))
        item["bank_profile"] = profile_name
        item["source_file"] = source_file
        item["bank_raw"] = raw
        item["bank_raw_description"] = str(description or "")
        item["bank_raw_amount"] = item.get("amount")
        item["bank_raw_date"] = item.get("date")
        converted.append(item)
    return converted


def import_table_with_profile(path: str | Path, profile: dict | None = None) -> tuple[list[dict], dict]:
    headers, rows = read_table_preview(path)
    detection = detect_profile(headers) if profile is None else {
        "profile": profile,
        "mapping": _profile_mapping(profile, headers),
        "confidence": 100,
        "ready": True,
    }
    if not detection["ready"]:
        raise ValueError("Nie rozpoznano układu kolumn. Użyj kreatora profilu banku.")
    selected = detection["profile"] or {"name": "Własny profil"}
    converted = convert_with_mapping(
        headers,
        rows,
        detection["mapping"],
        profile_name=selected.get("name", "Profil banku"),
        source_file=Path(path).name,
    )
    detection["transactions"] = len(converted)
    detection["headers"] = headers
    detection["preview"] = converted[:10]
    return converted, detection


def profile_summary(path: str | Path) -> dict:
    headers, rows = read_table_preview(path)
    detection = detect_profile(headers)
    detection["headers"] = headers
    detection["sample_rows"] = rows[:10]
    return detection
